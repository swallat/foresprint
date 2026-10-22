# SPDX-License-Identifier: EUPL-1.2
"""Golden references of the Excel import (``foresprint.excel_import``) for the Go port.

Workbooks are inputs: the export of every golden case, plus variants of the demo workbook that
are edited with openpyxl the way a user edits a sheet by hand (a wrong format version, text in a
number cell, a date as text, an unknown key, a missing sheet, ...). They are written below
``golden/expected/excel_import/workbooks`` with a fixed time stamp so that the bytes do not change
between runs. For each workbook the reference holds what ``read_workbook_keyed`` returns (the
documents it would write, as YAML) or the error message, then the plan against the existing
team folder (changes, files, untouched) and the bytes ``apply_import`` writes.
"""

from __future__ import annotations

import io
import re
import shutil
import tempfile
import zipfile
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.packaging.custom import StringProperty

from foresprint import excel_import
from foresprint.excel import build_workbook
from foresprint.forecast import build_forecast
from foresprint.storage import DataError, Repository, clear_team_cache, dump_yaml

FIXED = (2026, 1, 1, 0, 0, 0)
STAMP = "2026-01-01T00:00:00Z"


def normalise(raw: bytes) -> bytes:
    """The same workbook with fixed time stamps (zip entries and document properties)."""
    out = io.BytesIO()
    with (
        zipfile.ZipFile(io.BytesIO(raw)) as src,
        zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst,
    ):
        for info in src.infolist():
            data = src.read(info.filename)
            if info.filename == "docProps/core.xml":
                text = data.decode("utf-8")
                text = re.sub(r"(<dcterms:(?:created|modified)[^>]*>)[^<]*", rf"\g<1>{STAMP}", text)
                data = text.encode("utf-8")
            zi = zipfile.ZipInfo(info.filename, FIXED)
            zi.compress_type = zipfile.ZIP_DEFLATED
            zi.external_attr = info.external_attr
            dst.writestr(zi, data)
    return out.getvalue()


def edited(raw: bytes, edit) -> bytes:
    wb = load_workbook(io.BytesIO(raw))
    edit(wb)
    out = io.BytesIO()
    wb.save(out)
    return normalise(out.getvalue())


# --- editing helpers -------------------------------------------------------------------------
def setting(wb, key, value):
    ws = wb["Settings"]
    for row in range(3, ws.max_row + 1):
        if ws.cell(row, 3).value == key:
            ws.cell(row, 2).value = value
            return
    row = ws.max_row + 1
    ws.cell(row, 1).value = key
    ws.cell(row, 2).value = value
    ws.cell(row, 3).value = key


def forecast_row(wb, sid):
    ws = wb["Forecast"]
    for row in range(4, ws.max_row + 1):
        if ws.cell(row, 1).value == sid:
            return ws, row
    raise KeyError(sid)


def prop(wb, name, value):
    props = wb.custom_doc_props.props
    for p in props:
        if p.name == name:
            if value is None:
                props.remove(p)
            else:
                p.value = value
            return
    props.append(StringProperty(name=name, value=value))


def first_ids(wb):
    ws = wb["Forecast"]
    return [ws.cell(r, 1).value for r in range(4, ws.max_row + 1) if ws.cell(r, 1).value]


def variants() -> dict:
    """name -> edit(workbook); every edit starts from the export of the demo team."""
    v = {}

    def add(name):
        def deco(fn):
            v[name] = fn
            return fn

        return deco

    add("format-version-2")(lambda wb: prop(wb, "sprintforecast-format", "2"))
    add("no-properties")(
        lambda wb: (prop(wb, "sprintforecast-format", None), prop(wb, "sprintforecast-team", None))
    )
    add("other-team")(lambda wb: prop(wb, "sprintforecast-team", "other"))
    add("missing-sheet")(lambda wb: wb.remove(wb["People"]))

    def legacy(wb):
        for new, old in (
            ("Settings", "Einstellungen"), ("Absences", "Fehltage"), ("People", "Personen"),
            ("Sprint data", "Sprintdaten"), ("Product list", "Produktliste"),
            ("Product values", "Produktwerte"), ("Prognoses", "Prognosen"),
        ):  # fmt: skip
            wb[new].title = old

    add("legacy-sheet-names")(legacy)
    add("setting-text-for-int")(lambda wb: setting(wb, "forecast.window", "abc"))
    add("setting-float-for-int")(lambda wb: setting(wb, "forecast.window", 5.5))
    add("setting-int-as-float")(lambda wb: setting(wb, "forecast.window", 6.0))
    add("setting-decimal-comma")(lambda wb: setting(wb, "forecast.reference_capacity", "7,5"))
    add("setting-bad-flag")(lambda wb: setting(wb, "features.po", "vielleicht"))
    add("setting-flag-words")(
        lambda wb: (setting(wb, "features.po", "ja"), setting(wb, "features.production", "Nein"))
    )
    add("setting-unknown-key")(lambda wb: setting(wb, "foo.bar", 1))
    add("setting-label-without-dot")(lambda wb: setting(wb, "foo", 1))
    add("setting-empty-value")(lambda wb: setting(wb, "forecast.window", None))
    add("setting-bool-for-int")(lambda wb: setting(wb, "forecast.window", True))
    add("setting-huge")(lambda wb: setting(wb, "forecast.reference_capacity", 1e12))
    add("setting-date-text")(lambda wb: setting(wb, "sprint.first.start", "2026-01-05"))
    add("setting-date-bad-text")(lambda wb: setting(wb, "sprint.first.start", "05.01.2026"))
    add("setting-date-old")(lambda wb: setting(wb, "sprint.first.start", datetime(1980, 1, 1)))
    add("setting-window-zero")(lambda wb: setting(wb, "forecast.window", 0))
    add("setting-trend-mode-unknown")(lambda wb: setting(wb, "forecast.trend_mode", "wild"))
    add("setting-share-over-one")(lambda wb: setting(wb, "forecast.roadmap_share", 1.5))
    add("setting-name-number")(lambda wb: setting(wb, "name", 12.5))

    def set_id(wb, old, new):
        ws, row = forecast_row(wb, old)
        ws.cell(row, 1).value = new

    ids = "2026.1", "2026.2", "2026.3"
    add("id-float")(lambda wb: set_id(wb, ids[0], 2026.2))
    add("id-integer-float")(lambda wb: set_id(wb, ids[0], 2026.0))
    add("id-invalid")(lambda wb: set_id(wb, ids[1], "a:b"))
    add("id-case-duplicate")(lambda wb: (set_id(wb, ids[0], "Aa"), set_id(wb, ids[1], "aA")))

    def start(wb, value):
        ws, row = forecast_row(wb, ids[0])
        ws.cell(row, 2).value = value

    add("start-iso-text")(lambda wb: start(wb, "2026-01-12"))
    add("start-bad-text")(lambda wb: start(wb, "yesterday"))
    add("start-missing")(lambda wb: start(wb, None))
    add("start-time-only")(lambda wb: start(wb, datetime(2026, 1, 1).time()))
    add("start-1900")(lambda wb: start(wb, datetime(1900, 3, 1)))

    def number(wb, col, value):
        ws, row = forecast_row(wb, ids[0])
        ws.cell(row, col).value = value

    add("commitment-text")(lambda wb: number(wb, 4, "abc"))
    add("commitment-comma")(lambda wb: number(wb, 4, "11,5"))
    add("commitment-bool")(lambda wb: number(wb, 4, True))
    add("commitment-huge")(lambda wb: number(wb, 4, 5e9))
    add("commitment-spaces")(lambda wb: number(wb, 4, "  12  "))

    def free_text(wb):
        ws = wb["Forecast"]
        ws.cell(ws.max_row + 2, 1).value = "Sum"

    add("free-text-row")(free_text)

    def ghost(wb):
        ws = wb["Sprint data"]
        ws.cell(ws.max_row + 1, 1).value = "ghost"

    add("sprint-data-without-forecast-row")(ghost)

    def sd(wb, row, col, value):
        wb["Sprint data"].cell(row, col).value = value

    add("absence-state-bad")(lambda wb: sd(wb, 2, 2, "maybe"))
    add("absence-state-expected")(lambda wb: sd(wb, 2, 2, "expected"))
    add("corrections-bad")(lambda wb: sd(wb, 2, 6, "1 x"))
    add("corrections-comma")(lambda wb: sd(wb, 2, 6, "1,5 2"))
    add("planning-only-carryover")(
        lambda wb: (sd(wb, 2, 5, None), sd(wb, 2, 6, None), sd(wb, 2, 4, 3))
    )
    add("note-number")(lambda wb: sd(wb, 2, 11, 7.5))
    add("note-integer")(lambda wb: sd(wb, 2, 11, 7))

    def ab(wb, col, value):
        ws = wb["Absences"]
        ws.cell(2, col).value = value

    add("absence-unknown-kind")(lambda wb: ab(wb, 5, "foo"))
    add("absence-unknown-sprint")(lambda wb: ab(wb, 1, "nowhere"))
    add("absence-days-text")(lambda wb: ab(wb, 4, "x"))
    add("absence-days-missing")(lambda wb: ab(wb, 4, None))

    def pe(wb, row, col, value):
        wb["People"].cell(row, col).value = value

    add("person-fte-text")(lambda wb: pe(wb, 2, 9, "half"))
    add("person-country-lower")(lambda wb: pe(wb, 2, 3, "de"))
    add("person-access-word")(lambda wb: pe(wb, 2, 6, "ja"))
    add("person-access-bad")(lambda wb: pe(wb, 2, 6, "vielleicht"))
    add("person-no-from")(lambda wb: pe(wb, 2, 7, None))
    add("person-end-before-start")(lambda wb: pe(wb, 2, 8, datetime(2000, 1, 1)))

    def pr(wb, row, col, value):
        wb["Product list"].cell(row, col).value = value

    add("product-share-text")(lambda wb: pr(wb, 2, 6, "much"))
    add("product-no-name")(lambda wb: pr(wb, 2, 2, None))

    def pv(wb, sid, product, value, col=7):
        ws = wb["Product values"]
        row = ws.max_row + 1
        ws.cell(row, 1).value = sid
        ws.cell(row, 2).value = product
        ws.cell(row, col).value = value

    add("values-valid")(lambda wb: pv(wb, "2026.1", "webshop", 3.5))
    add("values-unknown-sprint")(lambda wb: pv(wb, "nowhere", "webshop", 3.5))
    add("values-text")(lambda wb: pv(wb, "2026.1", "webshop", "lots"))
    add("values-unknown-product")(lambda wb: pv(wb, "2026.1", "zzz", 3.5))

    def pg(wb, at=datetime(2026, 1, 8, 10, 30), sid="2026.1", preplanned=None, comment="plan"):
        ws = wb["Prognoses"]
        row = ws.max_row + 1
        ws.cell(row, 1).value = sid
        ws.cell(row, 2).value = at
        ws.cell(row, 3).value = comment
        ws.cell(row, 4).value = 12
        ws.cell(row, 8).value = "2026.4"
        ws.cell(row, 9).value = 10.5
        ws.cell(row, 10).value = preplanned

    add("prognosis-valid")(lambda wb: pg(wb, preplanned="webshop=1.5; app=2"))
    add("prognosis-two-unsorted")(
        lambda wb: (pg(wb, datetime(2026, 1, 9, 8, 0)), pg(wb, datetime(2026, 1, 8, 8, 0)))
    )
    add("prognosis-time-text")(lambda wb: pg(wb, "2026-01-01 10:00"))
    add("prognosis-date-only")(lambda wb: pg(wb, datetime(2026, 1, 8)))
    add("prognosis-preplanned-bad")(lambda wb: pg(wb, preplanned="A=x"))
    add("prognosis-preplanned-comma")(lambda wb: pg(wb, preplanned="A=1,5; B=2"))
    add("prognosis-preplanned-nan")(lambda wb: pg(wb, preplanned="A=nan"))
    add("prognosis-unknown-sprint")(lambda wb: pg(wb, sid="nowhere"))
    return v


def result_of(raw: bytes, key: str) -> dict:
    clear_team_cache()
    try:
        data, file_team = excel_import.read_workbook_keyed(raw, key)
        docs = {name: dump_yaml(doc) for name, doc in excel_import._documents(data).items()}
        return {"docs": docs, "file_team": file_team}
    except DataError as e:
        return {"error": str(e)}


def plan_of(raw: bytes, key: str, source: Path, today: date) -> dict:
    """The plan against a copy of the team folder ``source/key`` (or an empty folder)."""
    with tempfile.TemporaryDirectory() as tmp:
        root = Path(tmp)
        if (source / key).is_dir():
            shutil.copytree(source / key, root / key)
        clear_team_cache()
        repo = Repository(root)
        try:
            data, _ = excel_import.read_workbook_keyed(raw, key)
            plan = excel_import.plan_import(repo, data, today)
        except DataError as e:
            return {"error": str(e)}
        written = excel_import.apply_import(repo, plan)
        files = {
            str(p.relative_to(root / key)).replace("\\", "/"): p.read_text("utf-8") for p in written
        }
        return {
            "changes": [
                {"file": c.file, "path": c.path, "old": c.old, "new": c.new} for c in plan.changes
            ],
            "files": plan.files,
            "untouched": plan.untouched,
            "written": files,
        }


def generate(ctx) -> None:
    folder = ctx.expected / "excel_import" / "workbooks"
    folder.mkdir(parents=True, exist_ok=True)
    scenarios = []
    today_of = {}
    for case in ctx.resolved:
        data = ctx.load(case)
        raw = normalise(build_workbook(build_forecast(data, case.today)))
        name = f"case-{case.name}"
        (folder / f"{name}.xlsx").write_bytes(raw)
        scenarios.append(
            {
                "name": name,
                "key": case.team,
                "case": case.name,
                "root": case.root,
                "today": case.today,
            }
        )
        today_of[case.name] = case
    demo = next(c for c in ctx.resolved if c.name == "demo")
    base = normalise(build_workbook(build_forecast(ctx.load(demo), demo.today)))
    for name, edit in variants().items():
        raw = edited(base, edit)
        (folder / f"{name}.xlsx").write_bytes(raw)
        scenarios.append(
            {"name": name, "key": demo.team, "case": "demo", "root": demo.root, "today": demo.today}
        )
    # An older workbook without properties can be imported under a new key; a bad file is refused.
    old = edited(
        base,
        lambda wb: (prop(wb, "sprintforecast-format", None), prop(wb, "sprintforecast-team", None)),
    )
    (folder / "new-team.xlsx").write_bytes(old)
    scenarios.append(
        {"name": "new-team", "key": "brand-new", "case": None, "root": None, "today": demo.today}
    )
    (folder / "not-a-workbook.xlsx").write_bytes(b"this is not a zip file")
    scenarios.append(
        {
            "name": "not-a-workbook",
            "key": "demo",
            "case": "demo",
            "root": demo.root,
            "today": demo.today,
        }
    )
    (folder / "key-clash.xlsx").write_bytes(old)
    scenarios.append(
        {"name": "key-clash", "key": "DEMO", "case": "demo", "root": demo.root, "today": demo.today}
    )
    (folder / "key-bad.xlsx").write_bytes(old)
    scenarios.append(
        {"name": "key-bad", "key": "../x", "case": None, "root": None, "today": demo.today}
    )

    out = []
    for sc in scenarios:
        raw = (folder / f"{sc['name']}.xlsx").read_bytes()
        source = Path(sc["root"]) if sc["root"] else Path(tempfile.gettempdir())
        entry = {
            "name": sc["name"],
            "key": sc["key"],
            "case": sc["case"],
            "today": str(sc["today"]),
            "read": result_of(raw, sc["key"]),
        }
        if "docs" in entry["read"] or sc["name"] in ("key-clash", "key-bad", "new-team"):
            entry["plan"] = plan_of(
                raw,
                sc["key"],
                source,
                sc["today"] if isinstance(sc["today"], date) else date.fromisoformat(sc["today"]),
            )
        out.append(entry)
    ctx.write("excel_import/scenarios.json", ctx.dump_json(ctx.jsonable(out)))
