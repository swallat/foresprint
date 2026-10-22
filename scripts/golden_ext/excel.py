# SPDX-License-Identifier: EUPL-1.2
"""Golden references of the Excel output (``foresprint.excel``) for the Go port.

Bytes of a workbook differ between writers, so the reference is what the workbook says when it is
read back: for every sheet the cells (value, formula, cached result, number format, bold, fill,
font colour), column widths, row heights, frozen panes, conditional formats and charts, plus the
defined names and the custom properties. The Go test writes the same case and reads it back into
the same structure. Every golden case is exported in English, with one trend mode per case.
"""

from __future__ import annotations

import re
import zipfile
from datetime import date, datetime
from io import BytesIO
from xml.etree import ElementTree as ET

from openpyxl import load_workbook
from openpyxl.utils.datetime import to_excel

from foresprint.excel import build_workbook, import_unsafe_sprint_ids
from foresprint.forecast import build_forecast

NS = {
    "c": "http://schemas.openxmlformats.org/drawingml/2006/chart",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "xdr": "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing",
    "m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
}


def _plain(v):
    if isinstance(v, datetime):
        return float(to_excel(v))
    if isinstance(v, date):
        return float(to_excel(datetime(v.year, v.month, v.day)))
    if isinstance(v, bool):
        return v
    if isinstance(v, (int, float)):
        return float(v)
    return v


def _color(c):
    if c is None or c.type != "rgb" or not isinstance(c.rgb, str):
        return None
    return c.rgb[-6:].upper()


def _cells(ws_f, ws_v):
    cells = {}
    for row in ws_f.iter_rows():
        for cell in row:
            fill = (
                _color(cell.fill.fgColor) if cell.fill and cell.fill.fill_type == "solid" else None
            )
            font = cell.font
            style = {
                "nf": cell.number_format,
                "bold": bool(font.b),
                "size": float(font.sz) if font.sz else None,
                "fill": fill,
                "color": _color(font.color),
                "wrap": bool(cell.alignment.wrap_text),
                "valign": cell.alignment.vertical,
            }
            plain_style = {
                "nf": "General", "bold": False, "size": 11.0, "fill": None, "color": None,
                "wrap": False, "valign": None,
            }  # fmt: skip
            if cell.value is None and style == plain_style:
                continue
            entry = dict(style)
            v = cell.value
            if cell.data_type == "f":
                entry["f"] = v[1:] if isinstance(v, str) and v.startswith("=") else v
                cached = _plain(ws_v[cell.coordinate].value)
                kind = ws_v[cell.coordinate].data_type
                if cached is None:  # a formula that results in the empty text
                    cached, kind = "", "s"
                entry["c"], entry["ct"] = cached, "n" if kind == "d" else kind
            elif v is not None:
                entry["v"] = _plain(v)
                entry["t"] = {"s": "s", "b": "b"}.get(cell.data_type, "n")
            cells[cell.coordinate] = entry
    return cells


def _dims(ws):
    widths = {}
    for d in ws.column_dimensions.values():
        if d.width is None:
            continue
        for i in range(d.min or 0, (d.max or 0) + 1):
            widths[i] = round(float(d.width), 4)
    heights = {r: float(d.height) for r, d in ws.row_dimensions.items() if d.height is not None}
    return widths, heights


def _charts(z, sheet_path):
    rels = sheet_path.replace("worksheets/", "worksheets/_rels/") + ".rels"
    out = []
    if rels not in z.namelist():
        return out
    for rel in ET.fromstring(z.read(rels)):
        if not rel.get("Type", "").endswith("/drawing"):
            continue
        drawing = "xl/" + rel.get("Target").replace("../", "")
        d_rels = drawing.replace("drawings/", "drawings/_rels/") + ".rels"
        targets = {r.get("Id"): r.get("Target") for r in ET.fromstring(z.read(d_rels))}
        root = ET.fromstring(z.read(drawing))
        for anchor in root:
            frm = anchor.find("xdr:from", NS)
            to = anchor.find("xdr:to", NS)
            ident = anchor.find(".//c:chart", NS)
            if ident is None:
                continue
            rid = ident.get(
                "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
            )
            chart = ET.fromstring(z.read("xl/" + targets[rid].replace("../", "")))
            title = "".join(t.text or "" for t in chart.findall(".//c:chart/c:title//a:t", NS))
            series = []
            for ser in chart.findall(".//c:lineChart/c:ser", NS):
                series.append(
                    {
                        "name": ser.find("c:tx/c:strRef/c:f", NS).text,
                        "categories": ser.find("c:cat//c:f", NS).text,
                        "values": ser.find("c:val//c:f", NS).text,
                    }
                )
            axes = [
                "".join(t.text or "" for t in ax.findall("c:title//a:t", NS))
                for ax in chart.findall(".//c:catAx", NS) + chart.findall(".//c:valAx", NS)
            ]
            out.append(
                {
                    "title": title,
                    "series": series,
                    "axes": axes,
                    "from": [int(frm.find("xdr:col", NS).text), int(frm.find("xdr:row", NS).text)],
                    "to": [int(to.find("xdr:col", NS).text), int(to.find("xdr:row", NS).text)],
                    "legend": chart.find(".//c:legend/c:legendPos", NS).get("val"),
                    "na_blank": "dispNaAsBlank" in ET.tostring(chart, encoding="unicode"),
                    "blanks": (
                        chart.find(".//c:dispBlanksAs", NS) or ET.Element("x", val="gap")
                    ).get("val"),
                }
            )
    return out


def _conditional(sheet_xml):
    root = ET.fromstring(sheet_xml)
    out = []
    for cf in root.findall("m:conditionalFormatting", NS):
        for rule in cf.findall("m:cfRule", NS):
            out.append({"range": cf.get("sqref"), "type": rule.get("type")})
    return out


def dump(data: bytes) -> dict:
    z = zipfile.ZipFile(BytesIO(data))
    wb_f = load_workbook(BytesIO(data))
    wb_v = load_workbook(BytesIO(data), data_only=True)
    book = z.read("xl/workbook.xml").decode()
    paths = {}
    rels = {
        r.get("Id"): r.get("Target") for r in ET.fromstring(z.read("xl/_rels/workbook.xml.rels"))
    }
    for s in ET.fromstring(book).findall("m:sheets/m:sheet", NS):
        rid = s.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id")
        paths[s.get("name")] = "xl/" + rels[rid].lstrip("/").replace("xl/", "")
    sheets = []
    for ws in wb_f.worksheets:
        widths, heights = _dims(ws)
        sheets.append(
            {
                "name": ws.title,
                "freeze": ws.freeze_panes,
                "widths": {str(k): v for k, v in sorted(widths.items())},
                "heights": {str(k): v for k, v in sorted(heights.items())},
                "conditional": _conditional(z.read(paths[ws.title])),
                "charts": _charts(z, paths[ws.title]),
                "cells": _cells(ws, wb_v[ws.title]),
            }
        )
    names = {n: d.attr_text for n, d in wb_f.defined_names.items()}
    props = {p.name: p.value for p in wb_f.custom_doc_props.props}
    full_calc = re.search(r"<calcPr[^>]*fullCalcOnLoad=.1.", book) is not None
    return {"sheets": sheets, "names": names, "props": props, "full_calc": full_calc}


def generate(ctx) -> None:
    index = []
    for case in ctx.resolved:
        data = ctx.load(case)
        fc = build_forecast(data, case.today)
        raw = build_workbook(fc)
        ctx.write(f"excel/{case.name}.json", ctx.dump_json(dump(raw)))
        index.append({"case": case.name, "warnings": import_unsafe_sprint_ids(data)})
    ctx.write("excel/warnings.json", ctx.dump_json(index))
