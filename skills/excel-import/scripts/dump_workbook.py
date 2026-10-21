#!/usr/bin/env python3
"""Read-only dump of an Excel workbook for the excel-import skill.

Usage: dump_workbook.py WORKBOOK.xlsx [--max-rows N]
Prints each sheet (hidden ones flagged), its used range, merged cells and the non-empty cells
as `A1 = value` with the formula in brackets when there is one. Nothing is written.
"""

from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import load_workbook


def main(argv: list[str]) -> int:
    args = list(argv)
    max_rows = 500
    if "--max-rows" in args:
        i = args.index("--max-rows")
        if i + 1 >= len(args):
            print(__doc__)
            return 2
        max_rows = int(args[i + 1])
        del args[i : i + 2]
    if len(args) != 1:
        print(__doc__)
        return 2
    path = Path(args[0])
    values = load_workbook(path, data_only=True, read_only=False)
    formulas = load_workbook(path, data_only=False, read_only=False)
    for ws in values.worksheets:
        fs = formulas[ws.title]
        flag = "" if ws.sheet_state == "visible" else f"  [{ws.sheet_state.upper()}]"
        print(f"\n=== Sheet: {ws.title}{flag}  range {ws.dimensions} ===")
        if ws.merged_cells.ranges:
            print("merged:", ", ".join(str(r) for r in ws.merged_cells.ranges))
        for row in ws.iter_rows(max_row=max_rows):
            for cell in row:
                raw = fs[cell.coordinate].value
                if cell.value is None and raw is None:
                    continue
                note = f"   [formula {raw}]" if isinstance(raw, str) and raw.startswith("=") else ""
                print(f"{cell.coordinate} = {cell.value!r}{note}")
        if ws.max_row > max_rows:
            print(f"... {ws.max_row - max_rows} more rows (use --max-rows)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv[1:]))
