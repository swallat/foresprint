# SPDX-License-Identifier: EUPL-1.2
"""Excel-Ausgabe: echte Formeln, Zwischenspeicher aus der Engine, Diagramme, Befehl."""

import re
import zipfile
from datetime import date
from io import BytesIO
from pathlib import Path

import pytest

pytest.importorskip("xlsxwriter")

from typer.testing import CliRunner

from sprintforecast.cli import app
from sprintforecast.excel import ExcelUnavailable, build_workbook, trend_formula
from sprintforecast.forecast import build_forecast
from sprintforecast.model import TrendMode
from sprintforecast.storage import DATA_ENV, STICHTAG_ENV, Repository

TODAY = date(2026, 10, 15)


def sheet_xml(z: zipfile.ZipFile, name: str) -> str:
    """XML des Blatts ``name`` (Reihenfolge laut workbook.xml)."""
    book = z.read("xl/workbook.xml").decode()
    names = re.findall(r'<sheet name="([^"]+)"', book)
    return z.read(f"xl/worksheets/sheet{names.index(name) + 1}.xml").decode()


def workbook(data_dir: Path, mode: TrendMode = TrendMode.FLAT) -> zipfile.ZipFile:
    data = Repository(data_dir).load("demo")
    data.team.forecast.trend_mode = mode
    return zipfile.ZipFile(BytesIO(build_workbook(build_forecast(data, TODAY))))


def test_workbook_has_formulas_charts_and_the_settings_names(data_dir: Path):
    z = workbook(data_dir)
    names = z.namelist()
    assert sum(n.startswith("xl/charts/chart") for n in names) == 2  # Trend und Plan
    book = z.read("xl/workbook.xml").decode()
    for name in ("Referenz", "Fenster", "Roadmap_Anteil", "Rundung", "Duenn_nVelocity"):
        assert f'definedName name="{name}"' in book
    sheet = sheet_xml(z, "Forecast")
    formulas = re.findall(r"<f>(.*?)</f>", sheet)
    assert any("ROUNDDOWN" in f and "Referenz" in f for f in formulas)  # Forecast PLAN
    assert any("AVERAGEIFS" in f and "Fenster" in f for f in formulas)  # Durchschnitt
    assert any("Roadmap_Anteil" in f for f in formulas)
    assert len(formulas) > 200


def test_cached_values_come_from_the_engine(data_dir: Path):
    data = Repository(data_dir).load("demo")
    fc = build_forecast(data, TODAY)
    z = zipfile.ZipFile(BytesIO(build_workbook(fc)))
    sheet = sheet_xml(z, "Forecast")
    entry = fc.upcoming()[0]
    assert entry.result.plan is not None
    assert f"<v>{entry.result.plan}</v>" in sheet


@pytest.mark.parametrize("mode", list(TrendMode))
def test_trend_formula_exists_for_every_mode(mode: TrendMode):
    def ref(col: str, off: int) -> str:
        return f"{col}{10 - off}"

    text = trend_formula(mode, "Duenn", ref, 6)
    assert text.startswith("IF(Duenn,") and "known" in text


def test_products_sheet_and_production_split_are_included(data_dir: Path):
    z = workbook(data_dir)
    products = sheet_xml(z, "Produktwerte")
    assert "SUMIFS" not in products and "<f>" in products  # Differenz ohne Produktion
    forecast = sheet_xml(z, "Forecast")
    assert "SUMIFS" in forecast and "COUNTIFS" in forecast  # Σ Planning je Produkt


def test_imported_text_never_becomes_a_formula_in_the_next_export(data_dir: Path):
    """Notizen und Namen aus einer Mappe, die mit = beginnen, werden als Text geschrieben."""
    import openpyxl

    repo = Repository(data_dir)
    data = repo.load("demo")
    sprint = next(s for s in data.sprints if s.velocity is not None)
    sprint.note = '=HYPERLINK("http://boese.example","x")'
    data.people[0].name = "=1+1"
    repo.save_sprint("demo", sprint)
    repo.save_people("demo", data.people)
    content = build_workbook(build_forecast(repo.load("demo"), TODAY))
    wb = openpyxl.load_workbook(BytesIO(content))
    notes = [
        c for row in wb["Sprintdaten"].iter_rows(min_row=2) for c in row if c.value == sprint.note
    ]
    names = [c for row in wb["Personen"].iter_rows(min_row=2) for c in row if c.value == "=1+1"]
    assert notes and names
    assert all(c.data_type == "s" for c in [*notes, *names])  # Text, keine Formel ("f")


def test_command_writes_the_file(data_dir: Path, tmp_path: Path, monkeypatch: pytest.MonkeyPatch):
    monkeypatch.setenv(DATA_ENV, str(data_dir))
    monkeypatch.setenv(STICHTAG_ENV, TODAY.isoformat())
    target = tmp_path / "out.xlsx"
    result = CliRunner().invoke(app, ["excel", "demo", "-o", str(target)])
    assert result.exit_code == 0, result.output
    assert zipfile.is_zipfile(target)


def test_missing_xlsxwriter_is_reported(
    data_dir: Path, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
):
    import sprintforecast.excel as excel

    def missing() -> None:
        raise ExcelUnavailable("fehlt")

    monkeypatch.setattr(excel, "_xlsxwriter", missing)
    monkeypatch.setenv(DATA_ENV, str(data_dir))
    result = CliRunner().invoke(app, ["excel", "demo", "-o", str(tmp_path / "x.xlsx")])
    assert result.exit_code == 1 and not (tmp_path / "x.xlsx").exists()
