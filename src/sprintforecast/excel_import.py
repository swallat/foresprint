# SPDX-License-Identifier: EUPL-1.2
"""Import of an Excel workbook in the format of the Excel output (:mod:`sprintforecast.excel`):
the whole team including history: settings, people, products and all sprints with absences,
product values and prognoses.

Only the inputs (yellow cells) are read, by fixed column position; calculated columns
(formulas, standard distribution, names) are ignored. :func:`plan_import` first shows which
values change compared to the existing files; only :func:`apply_import` writes.
The import of the former plugin (:func:`sprintforecast.storage.import_sprints`) is not
affected. Requires ``openpyxl`` (``sprintforecast[excel]``)."""

from __future__ import annotations

import copy
import io
import math
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Any

import yaml

from sprintforecast.excel import (
    ABSENCE_KINDS,
    FORECAST_FIRST_ROW,
    FORMAT_PROPERTY,
    FORMAT_VERSION,
    LEGACY_SHEETS,
    SETTINGS_FIRST_ROW,
    SHEET_ABSENCES,
    SHEET_FORECAST,
    SHEET_PEOPLE,
    SHEET_PRODUCTS,
    SHEET_PROGNOSES,
    SHEET_SETTINGS,
    SHEET_SPRINTS,
    SHEET_VALUES,
    TEAM_PROPERTY,
    ExcelUnavailable,
)
from sprintforecast.i18n import t
from sprintforecast.model import (
    AbsenceState,
    Membership,
    Person,
    Product,
    ProductValues,
    Prognosis,
    Sprint,
    Team,
)
from sprintforecast.settings import SETTINGS
from sprintforecast.storage import (
    DataError,
    Repository,
    TeamData,
    _read,
    dump_yaml,
    file_name_problem,
    people_from_dict,
    people_to_dict,
    products_from_dict,
    products_to_dict,
    sprint_from_dict,
    sprint_to_dict,
    team_from_dict,
    team_key_problem,
    team_to_dict,
)
from sprintforecast.validate import Level, check_team


def _openpyxl() -> Any:
    try:
        import openpyxl
    except ImportError as e:
        raise ExcelUnavailable(
            t("The Excel import needs openpyxl: install sprintforecast with the [excel] extra.")
        ) from e
    return openpyxl


_TRUE = {"true", "wahr", "ja", "an", "yes", "on", "1", "x"}
_FALSE = {"false", "falsch", "nein", "aus", "no", "off", "0"}
# Plausible limits: quantities and days are far below them, dates fall in a sensible period.
MAX_NUMBER = 1e9
MIN_YEAR, MAX_YEAR = 1990, 2100


def check_sprint_id(sprint_id: str, where: str) -> str:
    """Sprint IDs become file names (see :func:`sprintforecast.storage.file_name_problem`)."""
    if problem := file_name_problem(sprint_id):
        raise DataError(
            t(
                "{where}: sprint ID {id} is not allowed ({problem})",
                where=where,
                id=repr(sprint_id),
                problem=problem,
            )
        )
    return sprint_id


class _Sheet:
    """A sheet with cell access and error messages that name the sheet and the cell."""

    def __init__(self, ws: Any) -> None:
        self.ws = ws
        self.name = ws.title

    def rows(self, first: int) -> list[tuple[int, tuple[Any, ...]]]:
        """Rows from ``first`` (1-based) to the last filled row."""
        out = []
        for number, row in enumerate(self.ws.iter_rows(min_row=first, values_only=True), first):
            if any(v not in (None, "") for v in row):
                out.append((number, row))
        return out

    def where(self, row: int, col: int) -> str:
        from openpyxl.utils import get_column_letter

        return f"{self.name}!{get_column_letter(col + 1)}{row}"

    def get(self, row: tuple[Any, ...], col: int) -> Any:
        value = row[col] if col < len(row) else None
        if isinstance(value, str):
            value = value.strip()
        return None if value == "" else value

    def number(self, number: int, row: tuple[Any, ...], col: int) -> float | None:
        value = self.get(row, col)
        if value is None:
            return None
        try:
            result = (
                float(str(value).replace(",", "."))
                if isinstance(value, bool) or not isinstance(value, int | float)
                else float(value)
            )
        except ValueError, OverflowError:  # text or an integer beyond float
            result = math.nan
        if not math.isfinite(result) or isinstance(value, bool):
            # Text, boolean, NaN or infinite: not a number that may be calculated with.
            raise DataError(
                t(
                    "{where}: a finite number expected, found: {value}",
                    where=self.where(number, col),
                    value=repr(value)[:60],
                )
            )
        if abs(result) > MAX_NUMBER:
            raise DataError(
                t(
                    "{where}: number outside the allowed range (at most {limit}): {value}",
                    where=self.where(number, col),
                    limit=f"{MAX_NUMBER:g}",
                    value=repr(value)[:60],
                )
            )
        return result

    def flag(self, number: int, row: tuple[Any, ...], col: int) -> bool:
        """Boolean value of a cell: Excel boolean or yes/no text; empty = no. Anything else is
        an error (``bool("aus")`` would be true)."""
        value = self.get(row, col)
        if value is None:
            return False
        if isinstance(value, bool):
            return value
        word = str(value).strip().lower()
        if word in _TRUE:
            return True
        if word in _FALSE:
            return False
        raise DataError(
            t(
                "{where}: yes or no (TRUE or FALSE) expected, found: {value}",
                where=self.where(number, col),
                value=repr(value),
            )
        )

    def text(self, number: int, row: tuple[Any, ...], col: int) -> str | None:
        value = self.get(row, col)
        if value is None:
            return None
        if (isinstance(value, float) and not value.is_integer()) or isinstance(value, datetime):
            raise DataError(
                t(
                    (
                        "{where}: text expected (enter numbers and dates as text, e.g. 2026.20 "
                        "in quotes or with a leading ' )"
                    ),
                    where=self.where(number, col),
                )
            )
        return str(int(value)) if isinstance(value, float) else str(value)

    def day(self, number: int, row: tuple[Any, ...], col: int) -> date | None:
        value = self.get(row, col)
        if value is None:
            return None
        if isinstance(value, datetime):
            day = value.date()
        elif isinstance(value, date):
            day = value
        else:
            try:
                day = date.fromisoformat(str(value))
            except ValueError:
                raise DataError(
                    t(
                        "{where}: date expected, found: {value}",
                        where=self.where(number, col),
                        value=repr(value)[:60],
                    )
                ) from None
        if not MIN_YEAR <= day.year <= MAX_YEAR:
            raise DataError(
                t(
                    "{where}: date outside {first} to {last}: {value}",
                    where=self.where(number, col),
                    first=MIN_YEAR,
                    last=MAX_YEAR,
                    value=day.isoformat(),
                )
            )
        return day


def _check_format(wb: Any) -> str | None:
    """Checks the format; returns the workbook's team key (None for older workbooks)."""
    for old, new in LEGACY_SHEETS.items():
        if old in wb.sheetnames and new not in wb.sheetnames:
            wb[old].title = new
    names = set(wb.sheetnames)
    needed = {SHEET_SETTINGS, SHEET_FORECAST, SHEET_SPRINTS, SHEET_PEOPLE, SHEET_PRODUCTS}
    if missing := sorted(needed - names):
        raise DataError(
            t(
                "Not an Excel output of sprintforecast: missing sheet(s): {missing}",
                missing=", ".join(missing),
            )
        )
    version = None
    file_team = None
    props = getattr(getattr(wb, "custom_doc_props", None), "props", None) or []
    for prop in props:
        if prop.name == FORMAT_PROPERTY:
            version = str(prop.value)
        elif prop.name == TEAM_PROPERTY:
            file_team = str(prop.value)
    if version is not None and version != FORMAT_VERSION:
        raise DataError(
            t(
                "The workbook has format {found}; this program reads format {expected}.",
                found=version,
                expected=FORMAT_VERSION,
            )
        )
    return file_team


_SETTING_INTS = {
    "sprint.length_days",
    "forecast.window",
    "forecast.rounding",
    "forecast.horizon",
    "capacity.rounding",
    "capacity.absence_history",
}
_SETTING_FLOATS = {
    "forecast.reference_capacity",
    "forecast.roadmap_share",
    "capacity.overhead_share",
    "capacity.overhead_days",
}
_SETTING_FLAGS = {"features.products", "features.po", "features.production", "features.tickets"}
_SETTING_DATES = {"sprint.first.start"}
_SETTING_TEXTS = {
    "name",
    "unit",
    "sprint.numbering",
    "sprint.first.id",
    "forecast.trend_mode",
    "forecast.product_basis",
    "capacity.country",
    "capacity.region",
    "capacity.workweek",
}


def _settings(sh: _Sheet) -> dict[str, Any]:
    """The settings as a nested dictionary as in ``team.yaml``. Only known keys, each value with
    its type: finite numbers (whole numbers as integers), yes/no switches."""
    result: dict[str, Any] = {}
    for number, row in sh.rows(SETTINGS_FIRST_ROW):
        path = sh.get(row, 2)
        if not isinstance(path, str) or not (
            path in _SETTING_TEXTS | _SETTING_INTS | _SETTING_FLOATS | _SETTING_FLAGS
            or path in _SETTING_DATES
        ):
            if isinstance(path, str) and path and sh.get(row, 1) is not None and "." in path:
                raise DataError(
                    t(
                        "{where}: unknown key {key}",
                        where=sh.where(number, 2),
                        key=repr(path),
                    )
                )
            continue  # label rows without a key
        if sh.get(row, 1) is None:
            continue
        value: Any
        if path in _SETTING_INTS:
            number_value = sh.number(number, row, 1)
            if number_value is None or not number_value.is_integer():
                raise DataError(
                    t(
                        "{where}: a whole number expected, found: {value}",
                        where=sh.where(number, 1),
                        value=repr(sh.get(row, 1)),
                    )
                )
            value = int(number_value)
        elif path in _SETTING_FLOATS:
            value = sh.number(number, row, 1)
        elif path in _SETTING_FLAGS:
            value = sh.flag(number, row, 1)
        elif path in _SETTING_DATES:
            value = sh.day(number, row, 1)
        else:
            value = sh.text(number, row, 1)
        node = result
        *parents, leaf = path.split(".")
        for key in parents:
            node = node.setdefault(key, {})
        node[leaf] = value
    return result


def _people(sh: _Sheet) -> list[Person]:
    people: dict[str, Person] = {}
    for number, row in sh.rows(2):
        pid = sh.text(number, row, 0)
        if pid is None:
            continue
        person = people.get(pid)
        if person is None:
            person = people[pid] = Person(
                id=pid,
                name=sh.text(number, row, 1) or pid,
                country=(sh.text(number, row, 2) or "").upper() or None,
                region=sh.text(number, row, 3),
                workweek=sh.text(number, row, 4),
                production_access=sh.flag(number, row, 5),
            )
        start = sh.day(number, row, 6)
        if start is not None:
            fte = sh.number(number, row, 8)
            person.memberships.append(
                Membership(start=start, end=sh.day(number, row, 7), fte=1.0 if fte is None else fte)
            )
    return list(people.values())


def _products(sh: _Sheet) -> list[Product]:
    products = []
    for number, row in sh.rows(2):
        pid = sh.text(number, row, 0)
        if pid is None:
            continue
        products.append(
            Product(
                id=pid,
                name=sh.text(number, row, 1) or pid,
                po=sh.text(number, row, 2),
                start=sh.day(number, row, 3),
                end=sh.day(number, row, 4),
                share=sh.number(number, row, 5),
            )
        )
    return products


def _preplanned(sh: _Sheet, number: int, row: tuple[Any, ...]) -> dict[str, float]:
    text = sh.text(number, row, 9)
    result: dict[str, float] = {}
    for part in (text or "").split(";"):
        if not part.strip():
            continue
        key, _, value = part.partition("=")
        try:
            number_value = float(value.strip().replace(",", "."))
            if not math.isfinite(number_value) or abs(number_value) > MAX_NUMBER:
                raise ValueError(value)
            result[key.strip()] = number_value
        except ValueError, OverflowError:
            raise DataError(
                t(
                    "{where}: expected ID=value; …, found: {value}",
                    where=sh.where(number, 9),
                    value=repr(text),
                )
            ) from None
    return result


def _sprints(wb: Any) -> list[Sprint]:
    fc = _Sheet(wb[SHEET_FORECAST])
    sp = _Sheet(wb[SHEET_SPRINTS])
    base: dict[str, dict[str, Any]] = {}
    seen_ids: set[str] = set()
    sheet_ids = {sp.text(n, r, 0) for n, r in sp.rows(2)} - {None}
    for number, row in fc.rows(FORECAST_FIRST_ROW + 1):
        sid = fc.text(number, row, 0)
        if sid is None:
            continue
        entry = {
            "start": fc.day(number, row, 1),
            "end": fc.day(number, row, 2),
            "forecast_ist": fc.number(number, row, 3),
            "velocity": fc.number(number, row, 4),
            "scope_creep": fc.number(number, row, 5),
            "capacity": fc.number(number, row, 6),
            "row": number,
        }
        if sid not in sheet_ids and all(entry[k] is None for k in entry if k != "row"):
            continue  # free text (sum, hint) below the sprints: not a sprint row
        check_sprint_id(sid, fc.where(number, 0))
        if sid.casefold() in seen_ids:
            raise DataError(
                t(
                    (
                        "{where}: sprint ID {id} appears twice (upper and lower case count as "
                        "the same on Windows and macOS)."
                    ),
                    where=fc.where(number, 0),
                    id=repr(sid),
                )
            )
        seen_ids.add(sid.casefold())
        base[sid] = entry
    listed: dict[str, tuple[int, tuple[Any, ...]]] = {}
    for number, row in sp.rows(2):
        if (sid := sp.text(number, row, 0)) is not None:
            check_sprint_id(sid, sp.where(number, 0))
            if sid.casefold() in {s.casefold() for s in listed}:
                raise DataError(
                    t(
                        (
                            "{where}: sprint ID {id} appears twice (upper and lower case count "
                            "as the same on Windows and macOS)."
                        ),
                        where=sp.where(number, 0),
                        id=repr(sid),
                    )
                )
            listed[sid] = (number, row)
    ids = [
        sid
        for sid, v in base.items()
        if sid in listed
        or any(v[k] is not None for k in ("forecast_ist", "velocity", "scope_creep", "capacity"))
    ]
    for sid in listed:
        if sid not in base:
            number = listed[sid][0]
            raise DataError(
                t(
                    "{where}: sprint {id} is missing on the sheet {sheet}",
                    where=sp.where(number, 0),
                    id=sid,
                    sheet=SHEET_FORECAST,
                )
            )
    sprints: dict[str, dict[str, Any]] = {}
    for sid in ids:
        v = base[sid]
        if v["start"] is None or v["end"] is None:
            raise DataError(
                t(
                    "{where}: start and end of sprint {id} are missing",
                    where=fc.where(v["row"], 1),
                    id=sid,
                )
            )
        d: dict[str, Any] = {
            "id": sid,
            "start": v["start"],
            "end": v["end"],
            "capacity": v["capacity"],
            "forecast_ist": v["forecast_ist"],
            "velocity": v["velocity"],
            "scope_creep": v["scope_creep"],
        }
        if sid in listed:
            number, row = listed[sid]
            if (state := sp.text(number, row, 1)) is not None:
                if state not in {s.value for s in AbsenceState}:
                    raise DataError(
                        t(
                            "{where}: unknown state of the absences: {value}",
                            where=sp.where(number, 1),
                            value=repr(state),
                        )
                    )
                d["absence_state"] = state
            for key, col in (
                ("forecast_plan", 2),
                ("not_done", 6),
                ("additional_done", 7),
                ("production_planned", 8),
                ("production_done", 9),
            ):
                d[key] = sp.number(number, row, col)
            carry, backlog = sp.number(number, row, 3), sp.number(number, row, 4)
            corr_text = sp.text(number, row, 5)
            if carry is not None or backlog is not None or corr_text:
                corrections = []
                for part in (corr_text or "").split():
                    try:
                        corrections.append(float(part.replace(",", ".")))
                    except ValueError:
                        raise DataError(
                            t(
                                (
                                    "{where}: please enter corrections as numbers separated by "
                                    "spaces"
                                ),
                                where=sp.where(number, 5),
                            )
                        ) from None
                d["planning"] = {
                    "carryover": carry or 0.0,
                    "backlog": backlog or 0.0,
                    "corrections": corrections,
                }
            d["note"] = sp.text(number, row, 10)
        sprints[sid] = d

    result = []
    for sid, d in sprints.items():
        raw = {k: v for k, v in d.items() if v is not None}
        result.append(sprint_from_dict(raw, f"{SHEET_FORECAST}: {sid}"))
    by_id = {s.id: s for s in result}

    ab = _Sheet(wb[SHEET_ABSENCES]) if SHEET_ABSENCES in wb.sheetnames else None
    for number, row in ab.rows(2) if ab else []:
        sid, pid, kind = ab.text(number, row, 0), ab.text(number, row, 1), ab.text(number, row, 4)
        days = ab.number(number, row, 3)
        if sid is None or pid is None or days is None:
            continue
        if sid not in by_id or kind not in ABSENCE_KINDS:
            raise DataError(
                t(
                    "{where}: unknown sprint {id} or kind {kind}",
                    where=ab.where(number, 0),
                    id=sid,
                    kind=repr(kind),
                )
            )
        getattr(by_id[sid], kind)[pid] = days

    vs = _Sheet(wb[SHEET_VALUES]) if SHEET_VALUES in wb.sheetnames else None
    measures = ("preplanned", "planned", "done", "production_planned", "production_done")
    for number, row in vs.rows(2) if vs else []:
        sid, pid = vs.text(number, row, 0), vs.text(number, row, 1)
        if sid is None or pid is None:
            continue
        values = ProductValues(**{m: vs.number(number, row, 5 + k) for k, m in enumerate(measures)})
        if values.is_empty():
            continue
        if sid not in by_id:
            raise DataError(t("{where}: unknown sprint {id}", where=vs.where(number, 0), id=sid))
        by_id[sid].products[pid] = values

    pg = _Sheet(wb[SHEET_PROGNOSES]) if SHEET_PROGNOSES in wb.sheetnames else None
    for number, row in pg.rows(2) if pg else []:
        sid = pg.text(number, row, 0)
        at = pg.get(row, 1)
        if sid is None or at is None:
            continue
        if sid not in by_id:
            raise DataError(t("{where}: unknown sprint {id}", where=pg.where(number, 0), id=sid))
        if not isinstance(at, datetime):
            raise DataError(
                t("{where}: point in time (date with time) expected", where=pg.where(number, 1))
            )
        by_id[sid].prognoses.append(
            Prognosis(
                at=at,
                comment=pg.text(number, row, 2),
                forecast_ist=pg.number(number, row, 3),
                not_done=pg.number(number, row, 4),
                additional_done=pg.number(number, row, 5),
                expected_done=pg.number(number, row, 6),
                target=pg.text(number, row, 7),
                plan=pg.number(number, row, 8),
                preplanned=_preplanned(pg, number, row),
            )
        )
    for sprint in result:  # the UI expects "older first"
        sprint.prognoses.sort(key=lambda p: p.at)
    return sorted(result, key=lambda s: s.start)


def read_workbook(source: Path | bytes, team_key: str) -> TeamData:
    """Reads the workbook into a :class:`TeamData`; errors name the sheet and the cell."""
    return read_workbook_keyed(source, team_key)[0]


def read_workbook_keyed(source: Path | bytes, team_key: str) -> tuple[TeamData, str | None]:
    """Like :func:`read_workbook`, but also returns the team key from the workbook.

    ``None`` means: older workbook without a key (callers should warn). If the workbook belongs
    to a different team, that is an error so that no team overwrites another team's data."""
    openpyxl = _openpyxl()
    stream = io.BytesIO(source) if isinstance(source, bytes) else source
    try:
        wb = openpyxl.load_workbook(stream, data_only=True)
    except Exception as e:  # not a valid xlsx file
        raise DataError(t("The file is not a readable Excel workbook: {error}", error=e)) from e
    file_team = _check_format(wb)
    if file_team is not None and file_team != team_key:
        raise DataError(
            t(
                (
                    "The workbook is from team {found}, but it is being imported into team "
                    "{expected}."
                ),
                found=repr(file_team),
                expected=repr(team_key),
            )
        )
    where = f"{SHEET_SETTINGS}"
    team = team_from_dict(team_key, _settings(_Sheet(wb[SHEET_SETTINGS])), where)
    data = TeamData(
        team=team,
        people=_people(_Sheet(wb[SHEET_PEOPLE])),
        products=_products(_Sheet(wb[SHEET_PRODUCTS])),
        sprints=_sprints(wb),
    )
    return data, file_team


@dataclass
class Change:
    file: str
    path: str
    old: Any
    new: Any


@dataclass
class ImportPlan:
    data: TeamData
    changes: list[Change] = field(default_factory=list)
    # Files that are written (relative to the team folder), and sprint files that already
    # exist but do not appear in the workbook (they stay unchanged).
    files: list[str] = field(default_factory=list)
    untouched: list[str] = field(default_factory=list)


def _documents(data: TeamData) -> dict[str, dict[str, Any]]:
    docs = {
        "team.yaml": team_to_dict(data.team),
        "people.yaml": people_to_dict(data.people),
        "products.yaml": products_to_dict(data.products),
    }
    for s in data.sprints:
        docs[f"sprints/{s.id}.yaml"] = sprint_to_dict(s)
    return docs


def _flat(doc: Any) -> dict[str, Any]:
    """Path → value; lists are broken down per entry with an index."""

    def walk(value: Any, prefix: str) -> dict[str, Any]:
        if isinstance(value, dict):
            out: dict[str, Any] = {}
            for k, v in value.items():
                out |= walk(v, f"{prefix}{k}.")
            return out
        if isinstance(value, list):
            out = {}
            for i, v in enumerate(value):
                out |= walk(v, f"{prefix}{i}.")
            return out
        return {prefix.rstrip("."): value}

    return walk(doc, "")


def check_team_key(repo: Repository, key: str) -> None:
    """The key is a plain folder name below the data folder (no path, no drive, no ``..``); a
    new team must also carry a permitted key and must not equal an existing one (not even up to
    upper and lower case)."""
    plain = (
        key not in ("", ".", "..")
        and key == Path(key).name
        and not any(c in key for c in "/\\:\x00")
        and not file_name_problem(key)
    )
    if not plain:
        raise DataError(t("Key {key}: not a plain folder name in the data folder.", key=repr(key)))
    if not repo.exists(key):
        if problem := team_key_problem(key):
            raise DataError(problem)
        if clash := repo.clashing_team(key):
            raise DataError(
                t("Team {team} already exists (as {existing}).", team=repr(key), existing=clash)
            )


def check_ids(data: TeamData) -> None:
    """Sprint IDs and the ID of the first sprint become file names: check them before they
    form a path."""
    for sprint in data.sprints:
        check_sprint_id(sprint.id, f"{SHEET_FORECAST}: {sprint.id}")
    if data.team.sprint.first_id is not None:
        check_sprint_id(data.team.sprint.first_id, f"{SHEET_SETTINGS}: sprint.first.id")


def _reload(data: TeamData) -> TeamData:
    """The data as they would be loaded again after writing: write each document as YAML and
    read it back with the same checks as when loading. Catches everything the own load check
    would reject later (non-finite numbers, wrong types, invalid IDs) before writing, and
    rounds like saving does (6 places)."""
    key = data.team.key

    def back(doc: dict[str, Any]) -> dict[str, Any]:
        return yaml.safe_load(dump_yaml(doc)) or {}

    team = team_from_dict(key, back(team_to_dict(data.team)), "team.yaml")
    people = people_from_dict(back(people_to_dict(data.people)), "people.yaml")
    products = products_from_dict(back(products_to_dict(data.products)), "products.yaml")
    sprints = [
        sprint_from_dict(back(sprint_to_dict(s)), f"sprints/{s.id}.yaml") for s in data.sprints
    ]
    return TeamData(team=team, people=people, sprints=sprints, products=products)


def check_settings(team: Team) -> None:
    """Every setting must pass the check of the settings page and of the command line (ranges,
    choices, region matching the country): ``Setting.parse`` on the displayed value."""
    problems = []
    for setting in SETTINGS:
        value = setting.get(team)
        if value is None and setting.optional:
            continue
        try:
            setting.parse(setting.display(team))
        except ValueError as e:
            problems.append(f"{setting.key}: {e}")
    if problems:
        raise DataError(
            t("The workbook contains invalid settings:")
            + "\n"
            + "\n".join(f"- {p}" for p in problems)
        )


def check_valid(data: TeamData, today: date, existing: TeamData | None = None) -> TeamData:
    """Rejects the workbook if, as it would be saved, it cannot be loaded or the data check
    (as before every save) finds errors. The overall state is checked: the workbook together
    with the existing sprint files that are missing from it. Returns the reloaded data."""
    check_ids(data)
    reloaded = _reload(data)
    check_settings(reloaded.team)
    union = reloaded
    if existing is not None:
        new_ids = {s.id for s in reloaded.sprints}
        extra = [s for s in existing.sprints if s.id not in new_ids]
        folded = {s.id.casefold(): s.id for s in reloaded.sprints}
        for old in extra:  # on Windows and macOS the same file as an imported ID
            if (clash := folded.get(old.id.casefold())) is not None:
                raise DataError(
                    t(
                        (
                            "Sprint ID {new} from the workbook and the existing sprint file "
                            "{old} differ only in upper and lower case (the same file on "
                            "Windows and macOS)."
                        ),
                        new=repr(clash),
                        old=repr(old.id),
                    )
                )
        union = TeamData(
            team=reloaded.team,
            people=reloaded.people,
            sprints=sorted([*reloaded.sprints, *extra], key=lambda s: s.start),
            products=reloaded.products,
        )
    errors = [i for i in check_team(union, today) if i.level is Level.ERROR]
    if errors:
        raise DataError(
            t("The workbook contains invalid values:")
            + "\n"
            + "\n".join(f"- {i.where}: {i.message}" for i in errors[:20])
        )
    return reloaded


def plan_import(repo: Repository, data: TeamData, today: date | None = None) -> ImportPlan:
    """Compares the workbook with the existing files: list of changes per file and value.
    Rejects an invalid key, invalid IDs, values that cannot be loaded and data with validation
    errors up front; nothing is written."""
    check_team_key(repo, data.team.key)
    existing = repo.load(data.team.key) if repo.exists(data.team.key) else None
    if existing is not None:
        # The workbook never contains the tracker section: the existing one is carried over.
        data.team.tracker = copy.deepcopy(existing.team.tracker)
        data.team.tracker_raw = copy.deepcopy(existing.team.tracker_raw)
    check_valid(data, today or date.today(), existing)
    team_dir = repo.team_dir(data.team.key)
    new_docs = _documents(data)
    plan = ImportPlan(data=data)
    for name, doc in new_docs.items():
        path = team_dir / name
        old = _read(path) if path.is_file() else None
        # Same representation as when writing: compare only after a round trip.
        old_flat = _flat(_normalize(name, old, data.team.key)) if old is not None else None
        new_flat = _flat(doc)
        if old_flat is None:
            plan.changes.append(Change(name, "", None, t("newly created")))
            plan.files.append(name)
            continue
        diff = False
        for key in sorted(old_flat.keys() | new_flat.keys()):
            before, after = old_flat.get(key), new_flat.get(key)
            if before != after:
                diff = True
                plan.changes.append(Change(name, key, before, after))
        if diff:
            plan.files.append(name)
    folder = team_dir / "sprints"
    if folder.is_dir():
        known = {f"sprints/{s.id}.yaml" for s in data.sprints}
        plan.untouched = sorted(
            f"sprints/{p.name}" for p in folder.glob("*.yaml") if f"sprints/{p.name}" not in known
        )
    return plan


def _normalize(name: str, raw: dict[str, Any], key: str) -> dict[str, Any]:
    """Bring an existing file into the same form in which the workbook would write it."""
    if name == "team.yaml":
        return team_to_dict(team_from_dict(key, raw, name))
    if name == "people.yaml":
        return people_to_dict(people_from_dict(raw, name))
    if name == "products.yaml":
        return products_to_dict(products_from_dict(raw, name))
    return sprint_to_dict(sprint_from_dict(raw, name))


def apply_import(repo: Repository, plan: ImportPlan) -> list[Path]:
    """Writes the changed files (unchanged ones are left untouched)."""
    data = plan.data
    key = data.team.key
    check_team_key(repo, key)  # here too: a plan can be assembled by hand
    check_ids(data)
    writers = {
        "team.yaml": lambda: repo.save_team(data.team),
        "people.yaml": lambda: repo.save_people(key, data.people),
        "products.yaml": lambda: repo.save_products(key, data.products),
    }
    paths = []
    sprints = {f"sprints/{s.id}.yaml": s for s in data.sprints}
    for name in plan.files:
        if name in writers:
            paths.append(writers[name]())
        else:
            paths.append(repo.save_sprint(key, sprints[name]))
    return paths
